# -*- coding: utf-8 -*-
"""把 Obsidian 暑期课程的总复习.md 转成带目录的 docx（逐行解析版 v3）。"""
import re, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

MD = r"C:\obsidian知识库\暑期课程\总复习.md"
OUT = r"C:\Users\LHW66\Documents\Codex\2026-08-10\chrome-plugin-chrome-openai-bundled-file\outputs\总复习.docx"
IMG_DIRS = [r"C:\obsidian知识库\图片库", r"C:\obsidian知识库\暑期课程"]

LATEX_REPL = {
    r"\pi": "π", r"\alpha": "α", r"\beta": "β", r"\theta": "θ", r"\varphi": "φ",
    r"\omega": "ω", r"\rho": "ρ", r"\mu": "μ", r"\lambda": "λ", r"\Delta": "Δ",
    r"\infty": "∞", r"\times": "×", r"\cdot": "·", r"\pm": "±", r"\mp": "∓",
    r"\ge": "≥", r"\le": "≤", r"\ne": "≠", r"\neq": "≠", r"\approx": "≈",
    r"\cup": "∪", r"\cap": "∩", r"\in": "∈", r"\subseteq": "⊆", r"\subset": "⊂",
    r"\Rightarrow": "⇒", r"\Leftrightarrow": "⇔", r"\to": "→", r"\rightarrow": "→",
    r"\Leftarrow": "⇐", r"\lvert": "|", r"\rvert": "|",
    r"\mathbb{R}": "ℝ", r"\mathbb{Z}": "ℤ", r"\mathbb{N}": "ℕ", r"\mathbb{Q}": "ℚ",
    r"\emptyset": "∅", r"\angle": "∠", r"\triangle": "△", r"\perp": "⊥",
    r"\parallel": "∥", r"\cdots": "…", r"\ldots": "…", r"\dots": "…",
    r"\sin": "sin", r"\cos": "cos", r"\tan": "tan", r"\cot": "cot",
    r"\ln": "ln", r"\lg": "lg", r"\log": "log",
    r"\dfrac": r"\frac", r"\cfrac": r"\frac", r"\tfrac": r"\frac",
}

def latex_to_text(s):
    s = s.replace("\\left", "").replace("\\right", "")
    s = re.sub(r"\\[,;!]|\\quad|\\qquad|\\ ", " ", s)
    s = re.sub(r"\\\\", " ", s)
    for k, v in LATEX_REPL.items():
        s = s.replace(k, v)
    s = re.sub(r"\\text\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\overrightarrow\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\vec\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\vec\s+([A-Za-z])", r"\1", s)
    s = re.sub(r"\\sqrt\{([^{}]*)\}", r"√(\1)", s)
    s = re.sub(r"\\sqrt\s*([0-9A-Za-z])", r"√\1", s)
    s = re.sub(r"\\frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)", s)
    s = re.sub(r"\\frac(\d)(\d)", r"(\1)/(\2)", s)
    s = re.sub(r"_\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"_([0-9A-Za-z\u4e00-\u9fff]+)", r"\1", s)
    SUP = {'0':"⁰",'1':"¹",'2':"²",'3':"³",'4':"⁴",'5':"⁵",'6':"⁶",'7':"⁷",'8':"⁸",'9':"⁹",'-':"⁻",'+':"⁺",'n':"ⁿ"}
    def sup_sub(m):
        inner = m.group(1)
        if all(c in SUP for c in inner):
            return "".join(SUP[c] for c in inner)
        return "^(" + inner + ")"
    s = re.sub(r"\^\{([^{}]*)\}", sup_sub, s)
    s = re.sub(r"\^([0-9A-Za-z-]+)", lambda m: "".join(SUP.get(c, c) for c in m.group(1)), s)
    s = re.sub(r"\\[a-zA-Z]+", "", s)
    s = s.replace("{", "").replace("}", "").replace("$", "")
    return s.strip()

WIKI_IMG = re.compile(r"!\[\[([^\]|]+?)(?:\|([^\]]*))?\]\]")
WIKI_LINK = re.compile(r"\[\[([^\]|]+?)(?:\|([^\]]*))?\]\]")

def find_image(name):
    base = os.path.basename(name)
    for d in IMG_DIRS:
        for root, _, files in os.walk(d):
            if base in files:
                return os.path.join(root, base)
    return None

def add_runs_with_markup(p, text, base_bold=False):
    def img_repl(m):
        name, size = m.group(1), m.group(2)
        low = name.lower()
        if low.endswith((".png", ".jpg", ".jpeg")):
            img = find_image(name)
            if img:
                try:
                    w = float(size) / 96.0 if size else 5.0
                    w = max(0.8, min(w, 6.0))
                    p.add_run().add_picture(img, width=Inches(w))
                except Exception:
                    r = p.add_run("[图片：%s]" % name); r.italic = True
                    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            else:
                r = p.add_run("[图片未找到：%s]" % name); r.italic = True
                r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        elif low.endswith(".html"):
            r = p.add_run("【互动演示：%s（HTML，请在 Obsidian 中打开）】" % os.path.basename(name))
            r.italic = True; r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        else:
            r = p.add_run(name); r.italic = True
        return ""
    def link_repl(m):
        p.add_run(m.group(1))
        return ""
    text = WIKI_IMG.sub(img_repl, text)
    text = WIKI_LINK.sub(link_repl, text)
    def math_repl(m):
        p.add_run(latex_to_text(m.group(1)))
        return ""
    text = re.sub(r"\$([^$\n]+)\$", math_repl, text)
    text = re.sub(r"<u>\s*</u>", "＿＿＿＿", text)
    text = re.sub(r"</?u>", "", text)
    text = re.sub(r"<sub>([^<]*)</sub>", r"\1", text)
    text = re.sub(r"<sup>([^<]*)</sup>", r"\1", text)
    parts = re.split(r"\*\*([^*]+)\*\*", text)
    for idx, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(part)
        if idx % 2 == 1 or base_bold:
            r.bold = True

def style_setup(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    for hname, sz in [("Heading 1", 17), ("Heading 2", 13.5), ("Heading 3", 12)]:
        st = doc.styles[hname]
        st.font.name = "Calibri"
        st.font.size = Pt(sz)
        st.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
        st.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    t = doc.styles["Title"]
    t.font.name = "Calibri"; t.font.size = Pt(26)
    t.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
    t.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")

lines = open(MD, encoding="utf-8").read().split("\n")
doc = Document()
style_setup(doc)

doc.add_paragraph("总复习（DAY1–DAY16）", style="Title")
sub = doc.add_paragraph()
r = sub.add_run("物理 DAY1–DAY8 ｜ 数学 DAY9–DAY16 ｜ 含全部答案")
r.italic = True; r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

doc.add_heading("目录", level=1)
for l in lines:
    m = re.match(r"^## (DAY\d+|数学部分|答案汇总)(.*)$", l.strip())
    if m:
        text = (m.group(1) + m.group(2)).strip()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(2)
        add_runs_with_markup(p, text)
doc.add_page_break()

i, n = 0, len(lines)
while i < n:
    s = lines[i].rstrip("\n").strip()
    if not s:
        i += 1
        continue
    if i == 0 and s.startswith("# "):
        i += 1
        continue
    if s.startswith("$$"):
        if s.endswith("$$") and len(s) > 4:
            body = s[2:-2].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(latex_to_text(body))
            r.italic = True
            i += 1
            continue
        buf = [s.strip("$ ").strip()]
        i += 1
        while i < n and not lines[i].strip().startswith("$$"):
            buf.append(lines[i].strip())
            i += 1
        if i < n:
            tail = lines[i].strip().lstrip("$$").strip()
            if tail:
                buf.append(tail)
            i += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(latex_to_text(" ".join(buf)))
        r.italic = True
        continue
    m = re.match(r"^(#{1,4})\s+(.*)$", s)
    if m:
        lv = len(m.group(1))
        h = doc.add_heading(m.group(2), level=lv - 1 if lv >= 2 else 1)
        i += 1
        continue
    if s.startswith("|"):
        rows = []
        while i < n and lines[i].strip().startswith("|"):
            rows.append(lines[i].strip())
            i += 1
        clean = [r for r in rows if not re.match(r"^\|[\s:\-|]+\|?$", r)]
        if clean:
            cells = [c.strip() for c in clean[0].strip("|").split("|")]
            table = doc.add_table(rows=1, cols=len(cells))
            table.style = "Table Grid"
            for j, c in enumerate(cells):
                add_runs_with_markup(table.rows[0].cells[j].paragraphs[0], c, base_bold=True)
            for row in clean[1:]:
                cells = [c.strip() for c in row.strip("|").split("|")]
                rc = table.add_row().cells
                for j in range(min(len(cells), len(rc))):
                    add_runs_with_markup(rc[j].paragraphs[0], cells[j])
        continue
    if re.match(r"^---+\s*$", s):
        i += 1
        continue
    if s.startswith(">"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(s.lstrip(">").strip())
        r.italic = True
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        i += 1
        continue
    m = re.match(r"^[-*]\s+(.*)$", s)
    if m:
        p = doc.add_paragraph(style="List Bullet")
        add_runs_with_markup(p, m.group(1))
        i += 1
        continue
    m = re.match(r"^(\d+)[\.、]\s*(.*)$", s)
    if m:
        p = doc.add_paragraph(style="List Number")
        add_runs_with_markup(p, m.group(2))
        i += 1
        continue
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_runs_with_markup(p, s)
    i += 1

doc.save(OUT)
print("saved:", os.path.getsize(OUT), "bytes")

