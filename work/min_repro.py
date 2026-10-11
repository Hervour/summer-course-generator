# -*- coding: utf-8 -*-
import sys, re
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
MD = r"C:\obsidian知识库\暑期课程\总复习.md"
lines = open(MD, encoding="utf-8").read().split("\n")
d = Document()
n_h1 = n_h2 = n_other = 0
for line in lines:
    s = line.rstrip("\n").strip()
    m = re.match(r"^(#{1,4})\s+(.*)$", s)
    if m:
        lv = len(m.group(1))
        d.add_heading(m.group(2), level=lv-1 if lv >= 2 else 1)
        if lv == 2: n_h1 += 1
        elif lv == 3: n_h2 += 1
        else: n_other += 1
d.save(r"C:\Users\LHW66\Documents\Codex\2026-08-10\chrome-plugin-chrome-openai-bundled-file\work\min.docx")
print("md 标题统计: ## =", n_h1, "### =", n_h2, "其他 =", n_other)
d2 = Document(r"C:\Users\LHW66\Documents\Codex\2026-08-10\chrome-plugin-chrome-openai-bundled-file\work\min.docx")
from collections import Counter
c = Counter(p.style.name for p in d2.paragraphs)
print("docx 样式统计:", dict(c))
h1s = [p.text for p in d2.paragraphs if p.style.name == "Heading 1"]
print("Heading 1 数量:", len(h1s))
print("缺:", [k for k in ["DAY11","DAY13","DAY15","DAY16 简单","答案汇总"] if not any(k in h for h in h1s)])
