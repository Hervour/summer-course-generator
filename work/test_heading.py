# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
d = Document()
h = d.add_heading("", level=1)
print("空标题 add_heading 样式:", repr(h.style.name))
r = h.add_run("DAY11 函数的基本概念")
print("加 run 后文本:", repr(h.text), "| 样式:", repr(h.style.name))
