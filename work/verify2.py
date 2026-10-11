# -*- coding: utf-8 -*-
from docx import Document
d = Document(r"C:\Users\LHW66\Documents\Codex\2026-08-10\chrome-plugin-chrome-openai-bundled-file\outputs\总复习.docx")
print("===== 全部 Heading 1 =====")
h1 = [p.text for p in d.paragraphs if p.style.name == "Heading 1" and p.text.strip()]
for h in h1:
    print("H1:", h[:60])
print()
print("===== 全部 Heading 2 =====")
h2 = [p.text for p in d.paragraphs if p.style.name == "Heading 2" and p.text.strip()]
for h in h2:
    print("H2:", h[:60])
