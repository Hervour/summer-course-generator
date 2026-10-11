# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
d = Document(r"C:\Users\LHW66\Documents\Codex\2026-08-10\chrome-plugin-chrome-openai-bundled-file\outputs\总复习.docx")
print("段落总数:", len(d.paragraphs), "| 图片数:", len(d.inline_shapes), "| 表格数:", len(d.tables))
h1 = [p.text for p in d.paragraphs if p.style.name == "Heading 1" and p.text.strip()]
print("Heading 1 总数:", len(h1))
print("===== Heading 1 =====")
for h in h1:
    print(" -", h[:55])
missing = [k for k in ["DAY11", "DAY13", "DAY15", "DAY16 简单", "答案汇总", "对数函数", "一、高中向量", "一、单选题"] if not any(k in h for h in h1)]
print()
print("缺失检查:", "无缺失" if not missing else missing)
