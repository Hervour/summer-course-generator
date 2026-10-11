# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
d = Document(r"C:\Users\LHW66\Documents\Codex\2026-08-10\chrome-plugin-chrome-openai-bundled-file\outputs\总复习.docx")
print("===== 含 DAY11 的段落 =====")
for p in d.paragraphs:
    if "DAY11" in p.text:
        print(f"[{p.style.name}] {p.text[:50]!r}")
print()
print("===== 含 答案汇总 的段落 =====")
for p in d.paragraphs:
    if "答案汇总" in p.text:
        print(f"[{p.style.name}] {p.text[:50]!r}")
