# -*- coding: utf-8 -*-
from docx import Document
d = Document(r"C:\Users\LHW66\Documents\Codex\2026-08-10\chrome-plugin-chrome-openai-bundled-file\outputs\总复习.docx")
paras = d.paragraphs
print("段落总数:", len(paras))
print("表格数:", len(d.tables))
print("图片数:", len(d.inline_shapes))
print()
print("===== 前 26 个非空段落（标题/目录区）=====")
cnt = 0
for p in paras:
    t = p.text.strip()
    if t:
        print(f"[{p.style.name}] {t[:58]}")
        cnt += 1
        if cnt >= 26:
            break
print()
print("===== Heading 1 数量与最后 6 个 =====")
h1 = [p.text for p in paras if p.style.name == "Heading 1" and p.text.strip()]
print("Heading 1 总数:", len(h1))
for h in h1[-6:]:
    print(" -", h[:60])
print()
print("===== 答案汇总章节标题 =====")
for p in paras:
    if p.style.name in ("Heading 1", "Heading 2") and ("答案" in p.text or "DAY8" in p.text or "DAY14" in p.text or "DAY16" in p.text):
        print(p.style.name, "|", p.text[:50])
